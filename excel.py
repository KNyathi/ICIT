# pip install openpyxl
import pandas as pd
import numpy as np
from itertools import combinations
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------- Data ----------
transactions = [
    ['ручка','тетрадь','карандаш'],
    ['ручка','тетрадь','ластик'],
    ['карандаш','линейка','ластик'],
    ['ручка','тетрадь','карандаш','ластик'],
    ['тетрадь','линейка'],
    ['ручка','карандаш'],
    ['ручка','тетрадь','линейка'],
    ['тетрадь','ластик','линейка'],
    ['ручка','тетрадь','карандаш'],
    ['карандаш','ластик'],
    ['ручка','тетрадь'],
    ['ручка','карандаш','линейка'],
    ['тетрадь','карандаш','ластик'],
    ['ручка','тетрадь','ластик'],
    ['ручка','линейка'],
    ['тетрадь','карандаш'],
    ['ручка','тетрадь','карандаш','линейка'],
    ['ластик','линейка'],
    ['ручка','тетрадь','карандаш'],
    ['тетрадь','ластик'],
    ['ручка','карандаш','ластик'],
    ['ручка','тетрадь','линейка'],
    ['карандаш','линейка'],
    ['ручка','тетрадь','карандаш','ластик'],
    ['тетрадь','линейка','ластик'],
    ['ручка','карандаш'],
    ['ручка','тетрадь','карандаш'],
    ['тетрадь','карандаш','линейка'],
    ['ручка','ластик'],
    ['ручка','тетрадь','карандаш'],
]
items = ['ручка','тетрадь','карандаш','ластик','линейка']
N = len(transactions)
min_sup = 0.3

# ---------- Binary matrix ----------
binary = pd.DataFrame(
    [[1 if it in t else 0 for it in items] for t in transactions],
    columns=items,
)
binary.index = [f"T{i+1}" for i in range(N)]
binary.index.name = "Transaction"

def support_count(itemset):
    mask = np.ones(N, dtype=bool)
    for it in itemset:
        mask &= (binary[it] == 1).values
    return int(mask.sum())

# ---------- Sheet 3: Item support ----------
item_rows = []
for it in items:
    c = int(binary[it].sum())
    item_rows.append({
        'Item': it,
        'Count': c,
        'Support': round(c / N, 3),
        'Formula': 'Count / N',
        'Calculation': f'{c} / {N} = {c/N:.3f}',
    })
item_support = pd.DataFrame(item_rows)

# ---------- Sheet 4: Pair support ----------
pair_rows = []
for a, b in combinations(items, 2):
    c = int(((binary[a] == 1) & (binary[b] == 1)).sum())
    pair_rows.append({
        'Item A': a, 'Item B': b,
        'Count': c,
        'Support': round(c / N, 3),
        'Formula': 'Σ(A=1 ∧ B=1) / N',
        'Calculation': f'{c} / {N} = {c/N:.3f}',
    })
pair_support = (pd.DataFrame(pair_rows)
                  .sort_values('Support', ascending=False)
                  .reset_index(drop=True))

# ---------- Sheet 5: Triple support ----------
triple_rows = []
for a, b, c in combinations(items, 3):
    cnt = int(((binary[a] == 1) & (binary[b] == 1) & (binary[c] == 1)).sum())
    triple_rows.append({
        'Itemset': f'{a}, {b}, {c}',
        'Count': cnt,
        'Support': round(cnt / N, 3),
        'Formula': 'Σ(A=1 ∧ B=1 ∧ C=1) / N',
        'Calculation': f'{cnt} / {N} = {cnt/N:.3f}',
    })
triple_support = (pd.DataFrame(triple_rows)
                    .sort_values('Support', ascending=False)
                    .reset_index(drop=True))

# ---------- Sheet 6: Frequent itemsets (min_support = 0.3) ----------
freq_rows = []
for it in items:
    c = int(binary[it].sum())
    if c / N >= min_sup:
        freq_rows.append({'Itemset': it, 'Size': 1, 'Count': c,
                          'Support': round(c/N, 3),
                          'Formula': 'Count / N',
                          'Calculation': f'{c} / {N} = {c/N:.3f} ≥ 0.300 ✓'})
for a, b in combinations(items, 2):
    c = int(((binary[a] == 1) & (binary[b] == 1)).sum())
    if c / N >= min_sup:
        freq_rows.append({'Itemset': f'{a}, {b}', 'Size': 2, 'Count': c,
                          'Support': round(c/N, 3),
                          'Formula': 'Σ(A=1 ∧ B=1) / N',
                          'Calculation': f'{c} / {N} = {c/N:.3f} ≥ 0.300 ✓'})
for a, b, cc in combinations(items, 3):
    c = int(((binary[a] == 1) & (binary[b] == 1) & (binary[cc] == 1)).sum())
    if c / N >= min_sup:
        freq_rows.append({'Itemset': f'{a}, {b}, {cc}', 'Size': 3, 'Count': c,
                          'Support': round(c/N, 3),
                          'Formula': 'Σ(A=1 ∧ B=1 ∧ C=1) / N',
                          'Calculation': f'{c} / {N} = {c/N:.3f} ≥ 0.300 ✓'})
frequent = (pd.DataFrame(freq_rows)
              .sort_values(['Size', 'Support'], ascending=[True, False])
              .reset_index(drop=True))

# ---------- Sheet 7: Rules ----------
RULE_COLUMNS = ['Antecedent', 'Consequent', 'Support', 'Confidence', 'Lift',
                'Support formula', 'Confidence formula', 'Lift formula',
                'Calculation']

def make_rules(freq_df, min_conf):
    rows = []
    for _, row in freq_df.iterrows():
        its = [x.strip() for x in row['Itemset'].split(',')]
        if len(its) < 2:
            continue
        for r in range(1, len(its)):
            for ant in combinations(its, r):
                con = tuple(x for x in its if x not in ant)
                c_both = support_count(its)
                c_ant  = support_count(ant)
                c_con  = support_count(con)
                sup_both = c_both / N
                conf = sup_both / (c_ant / N) if c_ant else 0
                lift = conf / (c_con / N) if c_con else 0
                if conf >= min_conf:
                    rows.append({
                        'Antecedent': ', '.join(ant),
                        'Consequent': ', '.join(con),
                        'Support': round(sup_both, 3),
                        'Confidence': round(conf, 3),
                        'Lift': round(lift, 3),
                        'Support formula': 'Count(A∪C) / N',
                        'Confidence formula': 'Support(A∪C) / Support(A)',
                        'Lift formula': 'Confidence / Support(C)',
                        'Calculation': (
                            f'sup={c_both}/{N}={sup_both:.3f}; '
                            f'conf={c_both}/{c_ant}={conf:.3f}; '
                            f'lift={conf:.3f}/{(c_con/N):.3f}={lift:.3f}'
                        ),
                    })
    if not rows:
        return pd.DataFrame(columns=RULE_COLUMNS)
    return (pd.DataFrame(rows, columns=RULE_COLUMNS)
              .drop_duplicates(subset=['Antecedent','Consequent'])
              .sort_values('Confidence', ascending=False)
              .reset_index(drop=True))

rules_60 = make_rules(frequent, 0.6)
rules_80 = make_rules(frequent, 0.8)

# ---------- Sheet 0: Constants ----------
constants = pd.DataFrame({
    'Parameter': ['N (transactions)', 'min_support', 'min_confidence (block 1)',
                  'min_confidence (block 2)', 'Number of items'],
    'Value': [N, min_sup, 0.6, 0.8, len(items)],
    'Meaning': [
        'Total number of transactions',
        'Minimum support threshold for frequent itemsets',
        'Threshold for rules block 1',
        'Threshold for rules block 2',
        'Distinct items in the store',
    ],
})

# ---------- Write everything ----------
out_path = 'variant6_association_rules.xlsx'
with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
    constants.to_excel(writer, sheet_name='0_Constants', index=False)

    tx_df = pd.DataFrame({
        'Transaction': [f'T{i+1}' for i in range(N)],
        'Items': [', '.join(t) for t in transactions],
        'Item count': [len(t) for t in transactions],
        'Formula': 'manual entry',
        'Calculation': ['count of items in the basket'] * N,
    })
    tx_df.to_excel(writer, sheet_name='1_Transactions', index=False)

    binary.reset_index().to_excel(writer, sheet_name='2_BinaryMatrix', index=False)

    item_support.to_excel(writer, sheet_name='3_ItemSupport', index=False)
    pair_support.to_excel(writer, sheet_name='4_PairSupport', index=False)
    triple_support.to_excel(writer, sheet_name='5_TripleSupport', index=False)
    frequent.to_excel(writer, sheet_name='6_FrequentItemsets', index=False)

    rules_60.to_excel(writer, sheet_name='7_Rules', index=False, startrow=0)
    header_row_80 = max(len(rules_60), 1) + 3
    rules_80.to_excel(writer, sheet_name='7_Rules', index=False,
                      startrow=header_row_80)
    ws = writer.sheets['7_Rules']
    ws.cell(row=header_row_80 - 1, column=1,
            value='Rules with Confidence ≥ 0.8').font = Font(bold=True)
    ws.cell(row=1, column=len(RULE_COLUMNS) + 2,
            value='Rules with Confidence ≥ 0.6').font = Font(bold=True)

# ---------- Styling ----------
wb = load_workbook(out_path)
header_fill = PatternFill('solid', fgColor='D9E1F2')
formula_fill = PatternFill('solid', fgColor='FFF2CC')   # soft yellow
bold = Font(bold=True)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    for cell in ws[1]:
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = center
    # Highlight Formula / Calculation columns
    headers = {c.value: c.column for c in ws[1] if c.value}
    for name in ('Formula', 'Support formula', 'Confidence formula',
                 'Lift formula', 'Calculation'):
        if name in headers:
            col = headers[name]
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=col).fill = formula_fill
    # Auto width
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 60)

wb.save(out_path)
print(f'Excel file saved: {out_path}')
print(f'Rules ≥ 0.6: {len(rules_60)}')
print(f'Rules ≥ 0.8: {len(rules_80)}')